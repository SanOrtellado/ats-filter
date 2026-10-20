"""Lógica del filtro ATS: extracción de texto, parsing, scoring y ranking de CVs.

Todo corre en local, sin APIs pagas:

* PyMuPDF / python-docx para leer PDF y DOCX.
* Regex para email, teléfono, LinkedIn, skills y años de experiencia.
* sentence-transformers (all-MiniLM-L6-v2) + cosine similarity de scikit-learn
  para el score semántico CV vs. Job Description.

El módulo también funciona como CLI para procesar una carpeta sin la interfaz::

    python ats_parser.py ./cvs --jd jd.txt --skills "python, sql, power bi" --out ranking.xlsx
"""

from __future__ import annotations

import argparse
import io
import logging
import os
import re
import sys
import unicodedata
from collections.abc import Callable, Iterable, Iterator, Sequence
from dataclasses import dataclass, field
from datetime import date
from functools import lru_cache
from pathlib import Path
from typing import TYPE_CHECKING, Any

import docx
import numpy as np
import pandas as pd
import pymupdf
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from sklearn.metrics.pairwise import cosine_similarity

if TYPE_CHECKING:
    from sentence_transformers import SentenceTransformer

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------

DEFAULT_MODEL_NAME = os.getenv("ATS_MODEL_NAME", "all-MiniLM-L6-v2")
SUPPORTED_EXTENSIONS: tuple[str, ...] = (".pdf", ".docx")

# Veredicto: PASA si score > 65, REVISAR si 45 <= score <= 65, NO PASA si < 45.
PASS_THRESHOLD = 65.0
REVIEW_THRESHOLD = 45.0
VERDICT_PASS = "PASA"
VERDICT_REVIEW = "REVISAR"
VERDICT_FAIL = "NO PASA"

# Peso del score semántico en el score final; el resto es cobertura de skills requeridas.
DEFAULT_SEMANTIC_WEIGHT = 0.7

# Detección de PDFs escaneados: promedio mínimo de caracteres útiles por página
# para considerar que el PDF tiene capa de texto, y umbral para marcar páginas sueltas.
MIN_CHARS_PER_PAGE = 100
SCANNED_PAGE_MAX_CHARS = 50

# all-MiniLM-L6-v2 trunca a 256 tokens: los textos se parten en ventanas solapadas.
CHUNK_WORDS = 120
CHUNK_OVERLAP = 30

# Calibración del coseno a 0-100. Con MiniLM, un CV de otro rubro contra una JD da
# ~0.15-0.30 y uno muy alineado ~0.65-0.85, así que [0.20, 0.80] se estira linealmente
# a [0, 100]. Ajustar con CVs reales si los veredictos salen corridos (ver README).
SIMILARITY_FLOOR = 0.20
SIMILARITY_CEILING = 0.80

MAX_REASONABLE_YEARS = 50.0

# ---------------------------------------------------------------------------
# Errores
# ---------------------------------------------------------------------------


class ExtractionError(Exception):
    """No se pudo obtener texto utilizable del archivo."""


class UnsupportedFileError(ExtractionError):
    """La extensión del archivo no está soportada."""


class ScannedPDFError(ExtractionError):
    """El PDF no tiene capa de texto (imagen escaneada): requiere OCR."""


# ---------------------------------------------------------------------------
# Modelos de datos
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Extraction:
    """Texto extraído de un CV más advertencias no bloqueantes."""

    text: str
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class ParsedCV:
    """Datos estructurados obtenidos del texto de un CV."""

    email: str | None
    phone: str | None
    linkedin: str | None
    skills: list[str]


@dataclass(frozen=True)
class CandidateResult:
    """Resultado completo de evaluar un CV contra una Job Description."""

    file_name: str
    email: str | None
    phone: str | None
    linkedin: str | None
    score: float
    semantic_score: float
    skill_coverage: float | None
    found_skills: list[str]
    missing_skills: list[str]
    experience_years: float | None
    verdict: str
    warnings: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class FileFailure:
    """Archivo que no pudo evaluarse, con el motivo."""

    file_name: str
    kind: str
    detail: str


# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------


def normalize_text(text: str) -> str:
    """Pasa a minúsculas y quita tildes (``Año`` -> ``ano``) para comparar con regex."""
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch)).lower()


def _clean_text(text: str) -> str:
    """Normaliza unicode (ligaduras, espacios duros) y colapsa espacios y saltos de línea."""
    text = unicodedata.normalize("NFKC", text).replace("\x00", "")
    lines = (" ".join(line.split()) for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _useful_chars(text: str) -> int:
    return len(re.sub(r"\s", "", text))


# ---------------------------------------------------------------------------
# Extracción de texto
# ---------------------------------------------------------------------------


def extract_text(file_path: str | Path) -> str:
    """Extrae el texto de un CV en PDF o DOCX.

    Args:
        file_path: Ruta al archivo ``.pdf`` o ``.docx``.

    Returns:
        Texto plano del CV, con los hipervínculos (mailto, LinkedIn) anexados al final.

    Raises:
        ScannedPDFError: Si el PDF es una imagen escaneada sin capa de texto.
        UnsupportedFileError: Si la extensión no es PDF ni DOCX.
        ExtractionError: Si el archivo no existe, está vacío, corrupto o protegido.
    """
    return extract_text_with_warnings(file_path).text


def extract_text_with_warnings(file_path: str | Path) -> Extraction:
    """Igual que :func:`extract_text`, pero devuelve también las advertencias no bloqueantes
    (por ejemplo, páginas sueltas escaneadas dentro de un PDF con texto)."""
    path = Path(file_path)
    if not path.is_file():
        raise ExtractionError(f"No existe el archivo: {path}")
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise ExtractionError(f"No se pudo leer el archivo: {exc}") from exc
    return extract_from_bytes(data, path.name)


def extract_from_bytes(data: bytes, file_name: str) -> Extraction:
    """Extrae texto a partir del contenido en memoria (por ejemplo, un archivo subido).

    Args:
        data: Bytes del archivo.
        file_name: Nombre original; se usa su extensión para elegir el lector.

    Raises:
        ScannedPDFError, UnsupportedFileError, ExtractionError: ver :func:`extract_text`.
    """
    suffix = Path(file_name).suffix.lower()
    if suffix == ".doc":
        raise UnsupportedFileError("Formato .doc (Word 97-2003) no soportado: convertir a .docx o PDF.")
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(f"Formato no soportado ({suffix or 'sin extensión'}): usar PDF o DOCX.")
    if not data:
        raise ExtractionError("El archivo está vacío.")

    extraction = _extract_pdf(data) if suffix == ".pdf" else _extract_docx(data)
    if not extraction.text.strip():
        raise ExtractionError("El archivo no contiene texto.")
    return extraction


def _extract_pdf(data: bytes) -> Extraction:
    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # PyMuPDF lanza distintos tipos según el daño del archivo
        raise ExtractionError(f"PDF corrupto o ilegible: {exc}") from exc

    with document:
        if document.needs_pass and not document.authenticate(""):
            raise ExtractionError("PDF protegido con contraseña.")
        if document.page_count == 0:
            raise ExtractionError("El PDF no tiene páginas.")

        pages: list[str] = []
        links: list[str] = []
        image_only_pages: list[int] = []
        has_images = False
        for page in document:
            page_text = page.get_text("text", sort=True)
            pages.append(page_text)
            links.extend(link["uri"] for link in page.get_links() if link.get("uri"))
            page_has_images = bool(page.get_images(full=False))
            has_images = has_images or page_has_images
            if page_has_images and _useful_chars(page_text) < SCANNED_PAGE_MAX_CHARS:
                image_only_pages.append(page.number + 1)
        page_count = document.page_count

    text = "\n".join(pages)
    if _useful_chars(text) < MIN_CHARS_PER_PAGE * page_count:
        if has_images:
            raise ScannedPDFError(
                "El PDF parece escaneado (imagen sin capa de texto). "
                "Pasarlo por OCR o pedir el CV en formato digital."
            )
        raise ExtractionError("El PDF no contiene texto extraíble.")

    warnings: list[str] = []
    if image_only_pages:
        pages_str = ", ".join(map(str, image_only_pages))
        warnings.append(f"Páginas sin texto, posiblemente escaneadas: {pages_str}. Su contenido no se evaluó.")
    return Extraction(_clean_text(_append_links(text, links)), tuple(warnings))


def _extract_docx(data: bytes) -> Extraction:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # zip inválido, XML roto, no es un .docx real, etc.
        raise ExtractionError(f"DOCX corrupto o ilegible: {exc}") from exc

    try:
        body = list(_iter_container_text(document))

        # Los encabezados suelen tener nombre y contacto; se deduplican porque las
        # secciones pueden heredar el mismo encabezado.
        header_footer: list[str] = []
        for section in document.sections:
            for part in (section.header, section.footer):
                header_footer.extend(_iter_container_text(part))

        # python-docx no expone los cuadros de texto (muy usados en CVs "de diseño").
        text_boxes: list[str] = []
        for box in document.element.body.iter(qn("w:txbxContent")):
            for paragraph in box.iter(qn("w:p")):
                text_boxes.append("".join(node.text or "" for node in paragraph.iter(qn("w:t"))))

        links = [
            rel.target_ref
            for rel in document.part.rels.values()
            if rel.reltype == RT.HYPERLINK and rel.is_external
        ]
    except Exception as exc:
        raise ExtractionError(f"No se pudo recorrer el contenido del DOCX: {exc}") from exc

    parts = [*dict.fromkeys(header_footer), *body, *dict.fromkeys(text_boxes)]
    return Extraction(_clean_text(_append_links("\n".join(parts), links)))


def _iter_container_text(container: Any) -> Iterator[str]:
    """Recorre párrafos y tablas (incluidas tablas anidadas) en orden de documento."""
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            yield block.text
            continue
        for row in block.rows:
            seen: set[int] = set()
            for cell in row.cells:
                # Las celdas combinadas se repiten en row.cells: se procesan una sola vez.
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                yield from _iter_container_text(cell)


def _append_links(text: str, links: Iterable[str]) -> str:
    """Anexa URLs de hipervínculos: a veces el LinkedIn o el mail solo están en el link."""
    unique = [link for link in dict.fromkeys(links) if link and link not in text]
    return f"{text}\n{chr(10).join(unique)}" if unique else text


def iter_cv_files(folder: str | Path, recursive: bool = False) -> list[Path]:
    """Lista los CVs (PDF/DOCX) de una carpeta, ordenados por nombre.

    Ignora archivos temporales de Office (``~$archivo.docx``) y ocultos.

    Raises:
        NotADirectoryError: Si la ruta no es una carpeta existente.
    """
    root = Path(folder).expanduser()
    if not root.is_dir():
        raise NotADirectoryError(f"La carpeta no existe: {root}")
    pattern = "**/*" if recursive else "*"
    return sorted(
        (
            path
            for path in root.glob(pattern)
            if path.is_file()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
            and not path.name.startswith(("~$", "."))
        ),
        key=lambda p: str(p).lower(),
    )


# ---------------------------------------------------------------------------
# Skills
# ---------------------------------------------------------------------------

# Nombre canónico -> alias como fragmentos regex sobre texto normalizado (minúsculas,
# sin tildes). Cada nombre canónico debe matchear alguno de sus propios alias.
SKILL_CATALOG: dict[str, tuple[str, ...]] = {
    # Lenguajes
    "python": (r"python",),
    "pyspark": (r"pyspark",),
    "r": (r"rstudio", r"r studio", r"r", r"lenguaje r", r"programacion en r", r"r programming", r"tidyverse", r"ggplot2?"),
    "sql": (r"sql", r"t-?sql", r"pl/?sql", r"spark sql"),
    "scala": (r"scala",),
    "java": (r"java(?! ?script)",),
    "javascript": (r"javascript", r"java script", r"ecmascript"),
    "typescript": (r"typescript",),
    "c#": (r"c#", r"c sharp", r"csharp"),
    "c++": (r"c\+\+", r"cpp"),
    ".net": (r"\.net", r"asp\.net", r"dotnet"),
    "go": (r"golang", r"go"),
    "bash": (r"bash", r"shell scripting"),
    "vba": (r"vba", r"visual basic for applications", r"macros? (?:de|en) excel", r"excel macros?"),
    "dax": (r"dax",),
    "power query": (r"power ?query", r"lenguaje m"),
    # Ecosistema Microsoft
    "power bi": (r"power[ -]?bi",),
    "fabric": (r"fabric", r"microsoft fabric", r"ms fabric"),
    "azure": (r"azure",),
    "azure data factory": (r"azure data factory", r"data factory", r"adf"),
    "synapse": (r"synapse", r"azure synapse"),
    "sql server": (r"sql server", r"ms ?sql", r"mssql"),
    "ssis": (r"ssis",),
    "ssas": (r"ssas",),
    "ssrs": (r"ssrs",),
    "excel": (r"excel",),
    "power automate": (r"power ?automate", r"microsoft flow"),
    "power apps": (r"power ?apps?",),
    # Ingeniería de datos
    "etl": (r"etl", r"elt", r"extraccion,? transformacion y carga", r"extract,? transform,? (?:and|&) load"),
    "spark": (r"spark", r"apache spark"),
    "databricks": (r"databricks",),
    "hadoop": (r"hadoop", r"hdfs"),
    "hive": (r"hive",),
    "kafka": (r"kafka",),
    "airflow": (r"airflow",),
    "dbt": (r"dbt",),
    "snowflake": (r"snowflake",),
    "bigquery": (r"big ?query",),
    "redshift": (r"redshift",),
    "data warehouse": (r"data ?warehouse", r"data ?warehousing", r"dwh", r"almacen(?:es)? de datos"),
    "data lake": (r"data ?lake", r"data ?lakehouse", r"lakehouse"),
    "modelado de datos": (
        r"modelado (?:de datos|dimensional)",
        r"data modell?ing",
        r"dimensional modell?ing",
        r"star schema",
        r"esquema (?:en )?estrella",
        r"kimball",
    ),
    # Cloud
    "aws": (r"aws", r"amazon web services"),
    "gcp": (r"gcp", r"google cloud(?: platform)?"),
    # Bases de datos
    "postgresql": (r"postgre ?sql", r"postgres"),
    "mysql": (r"mysql",),
    "oracle": (r"oracle",),
    "mongodb": (r"mongo ?db", r"mongo"),
    "nosql": (r"nosql",),
    # BI y analítica
    "tableau": (r"tableau",),
    "looker": (r"looker(?! ?studio)",),
    "looker studio": (r"looker ?studio", r"(?:google )?data ?studio"),
    "qlik": (r"qlik ?(?:view|sense)?",),
    "google sheets": (r"google sheets", r"hojas de calculo de google"),
    "google analytics": (r"google analytics", r"ga4", r"universal analytics"),
    "estadística": (r"estadistica", r"statistics", r"statistical analysis", r"analisis estadistico"),
    "machine learning": (r"machine learning", r"aprendizaje automatico"),
    "deep learning": (r"deep learning", r"aprendizaje profundo"),
    "nlp": (r"nlp", r"procesamiento de lenguaje natural", r"natural language processing"),
    "pandas": (r"pandas",),
    "numpy": (r"numpy",),
    "scikit-learn": (r"scikit[ -]?learn", r"sklearn"),
    "tensorflow": (r"tensorflow",),
    "pytorch": (r"pytorch",),
    # Ingeniería de software / DevOps
    "git": (r"git", r"github", r"gitlab", r"bitbucket"),
    "docker": (r"docker",),
    "kubernetes": (r"kubernetes", r"k8s"),
    "linux": (r"linux",),
    "ci/cd": (r"ci ?/ ?cd", r"cicd", r"integracion continua", r"continuous integration"),
    "api rest": (r"apis? rest(?:ful)?", r"rest(?:ful)? apis?"),
    "sap": (r"sap",),
    "salesforce": (r"salesforce",),
    # Metodologías e idiomas
    "scrum": (r"scrum",),
    "agile": (r"agile", r"metodologias? agil(?:es)?", r"kanban"),
    "jira": (r"jira",),
    "inglés": (r"ingles", r"english"),
}

# Alias que se aceptan al canonizar lo que escribe el usuario, pero que son demasiado
# ambiguos para buscarlos sueltos en un CV ("R", "go").
_INPUT_ONLY_ALIASES: dict[str, frozenset[str]] = {"r": frozenset({r"r"}), "go": frozenset({r"go"})}

# Una skill detectada implica otras (quien usa PostgreSQL sabe SQL, PySpark implica Spark y Python).
IMPLIED_SKILLS: dict[str, tuple[str, ...]] = {
    "pyspark": ("spark", "python"),
    "postgresql": ("sql",),
    "mysql": ("sql",),
    "sql server": ("sql",),
    "databricks": ("spark",),
    "azure data factory": ("azure", "etl"),
    "synapse": ("azure",),
    "ssis": ("etl",),
    "pandas": ("python",),
    "numpy": ("python",),
    "scikit-learn": ("python",),
}

_BOUNDARY_BEFORE = r"(?<![a-z0-9_+#])"
_BOUNDARY_AFTER = r"(?![a-z0-9_+#])"


def _alternation(aliases: Iterable[str]) -> str:
    return "(?:" + "|".join(aliases) + ")"


def _compile_search(aliases: Iterable[str]) -> re.Pattern[str]:
    return re.compile(_BOUNDARY_BEFORE + _alternation(aliases) + _BOUNDARY_AFTER)


_SKILL_SEARCH: dict[str, re.Pattern[str]] = {
    skill: _compile_search(a for a in aliases if a not in _INPUT_ONLY_ALIASES.get(skill, frozenset()))
    for skill, aliases in SKILL_CATALOG.items()
}
_SKILL_FULLMATCH: dict[str, re.Pattern[str]] = {
    skill: re.compile(_alternation(aliases)) for skill, aliases in SKILL_CATALOG.items()
}


def canonicalize_skill(skill: str) -> str:
    """Devuelve el nombre canónico de una skill (``"PowerBI"`` -> ``"power bi"``).

    Si no está en :data:`SKILL_CATALOG`, devuelve el texto en minúsculas tal cual,
    y se buscará literalmente en los CVs.
    """
    display = " ".join(skill.lower().split())
    normalized = normalize_text(display)
    for canonical, pattern in _SKILL_FULLMATCH.items():
        if pattern.fullmatch(normalized):
            return canonical
    return display


def parse_skill_list(raw: str) -> list[str]:
    """Convierte ``"Python, SQL; Power BI"`` en una lista canónica sin duplicados."""
    items = (canonicalize_skill(item) for item in re.split(r"[,;\n]", raw or "") if item.strip())
    return list(dict.fromkeys(items))


def _skill_pattern(skill: str) -> re.Pattern[str]:
    if skill in _SKILL_SEARCH:
        return _SKILL_SEARCH[skill]
    words = normalize_text(skill).split()
    return _compile_search([r"\s+".join(re.escape(word) for word in words)])


def detect_skills(text: str, extra_skills: Iterable[str] = ()) -> list[str]:
    """Detecta skills del catálogo (y opcionalmente otras a medida) en el texto de un CV.

    Args:
        text: Texto del CV.
        extra_skills: Skills adicionales a buscar aunque no estén en el catálogo
            (típicamente las requeridas por la búsqueda).

    Returns:
        Skills canónicas encontradas, en orden alfabético.
    """
    normalized = normalize_text(text)
    found = {skill for skill, pattern in _SKILL_SEARCH.items() if pattern.search(normalized)}
    for skill in list(found):
        found.update(IMPLIED_SKILLS.get(skill, ()))
    for extra in extra_skills:
        canonical = canonicalize_skill(extra)
        if canonical and canonical not in found and _skill_pattern(canonical).search(normalized):
            found.add(canonical)
    return sorted(found)


def get_missing_skills(found_skills: Iterable[str], required_skills: Iterable[str]) -> list[str]:
    """Devuelve las skills requeridas que no aparecen en el CV, en el orden en que se pidieron.

    Ambas listas se canonizan, así que ``"PowerBI"`` y ``"power bi"`` se consideran iguales.
    """
    found = {canonicalize_skill(skill) for skill in found_skills}
    required = dict.fromkeys(canonicalize_skill(skill) for skill in required_skills if skill.strip())
    return [skill for skill in required if skill not in found]


def skill_coverage(found_skills: Iterable[str], required_skills: Sequence[str]) -> float | None:
    """Porcentaje (0-100) de skills requeridas presentes; ``None`` si no hay requeridas."""
    required = list(dict.fromkeys(canonicalize_skill(skill) for skill in required_skills if skill.strip()))
    if not required:
        return None
    missing = get_missing_skills(found_skills, required)
    return round(100.0 * (len(required) - len(missing)) / len(required), 1)


# ---------------------------------------------------------------------------
# Parsing de datos de contacto
# ---------------------------------------------------------------------------

_EMAIL_RE = re.compile(r"[a-z0-9._%+-]+@[a-z0-9-]+(?:\.[a-z0-9-]+)*\.[a-z]{2,}")
_LINKEDIN_RE = re.compile(
    r"(?:https?://)?(?:[a-z]{2,3}\.)?linkedin\.com/(in|pub)/([\w\-%]+)", re.IGNORECASE
)
_PHONE_CANDIDATE_RE = re.compile(r"(?<![\w+])\+?\(?\d[\d \t().-]{6,22}\d(?!\w)")
_PHONE_LABEL_RE = re.compile(r"\b(?:tel|telefono|phone|cel|celular|movil|mobile|whatsapp|wpp|wsp)\b")
_ID_LABEL_RE = re.compile(r"\b(?:dni|cuit|cuil|cedula|pasaporte|passport|documento|rut|nif|nie|ssn|id)\b")
_DATE_LIKE_RE = re.compile(r"^\d{1,2}[./-]\d{1,2}[./-](?:19|20)\d{2}$")


def parse_cv(text: str, extra_skills: Iterable[str] = ()) -> ParsedCV:
    """Extrae email, teléfono, LinkedIn y skills del texto de un CV.

    Args:
        text: Texto del CV (salida de :func:`extract_text`).
        extra_skills: Skills a buscar además del catálogo predefinido.

    Returns:
        :class:`ParsedCV` con los campos encontrados (``None`` si no aparece).
    """
    normalized = normalize_text(text)

    email_match = _EMAIL_RE.search(normalized)
    email = email_match.group().strip(".") if email_match else None

    linkedin_match = _LINKEDIN_RE.search(text)
    linkedin = (
        f"https://www.linkedin.com/{linkedin_match.group(1).lower()}/{linkedin_match.group(2)}"
        if linkedin_match
        else None
    )

    return ParsedCV(
        email=email,
        phone=_extract_phone(normalized),
        linkedin=linkedin,
        skills=detect_skills(text, extra_skills),
    )


def _extract_phone(normalized_text: str) -> str | None:
    """Prefiere números precedidos por una etiqueta de teléfono (Tel, Cel, WhatsApp...)
    y descarta los que siguen a una etiqueta de documento (DNI, CUIT, pasaporte...)."""
    fallback: str | None = None
    for line in normalized_text.splitlines():
        labels = sorted(
            [(m.end(), "phone") for m in _PHONE_LABEL_RE.finditer(line)]
            + [(m.end(), "id") for m in _ID_LABEL_RE.finditer(line)]
        )
        for match in _PHONE_CANDIDATE_RE.finditer(line):
            candidate = " ".join(match.group().split()).strip(" .-")
            if not _looks_like_phone(candidate):
                continue
            preceding = [kind for end, kind in labels if end <= match.start()]
            if preceding and preceding[-1] == "phone":
                return candidate
            if not preceding or preceding[-1] != "id":
                fallback = fallback or candidate
    return fallback


def _looks_like_phone(candidate: str) -> bool:
    digits = re.sub(r"\D", "", candidate)
    if not 8 <= len(digits) <= 15 or _DATE_LIKE_RE.match(candidate):
        return False
    groups = re.findall(r"\d+", candidate)
    # "2018 - 2023" o "2015 2019" son rangos de años, no teléfonos.
    return not all(len(g) == 4 and 1950 <= int(g) <= 2049 for g in groups)


# ---------------------------------------------------------------------------
# Años de experiencia
# ---------------------------------------------------------------------------

_NUMBER_WORDS: dict[str, int] = {
    "un": 1, "uno": 1, "one": 1, "dos": 2, "two": 2, "tres": 3, "three": 3,
    "cuatro": 4, "four": 4, "cinco": 5, "five": 5, "seis": 6, "six": 6,
    "siete": 7, "seven": 7, "ocho": 8, "eight": 8, "nueve": 9, "nine": 9,
    "diez": 10, "ten": 10, "once": 11, "eleven": 11, "doce": 12, "twelve": 12,
    "quince": 15, "fifteen": 15, "veinte": 20, "twenty": 20,
}
_NUM = (
    r"(?<![\w.,])(?P<n>\d{1,2}(?:[.,]\d)?|"
    + "|".join(sorted(_NUMBER_WORDS, key=len, reverse=True))
    + r")(?![\w])"
)
_YEARS = r"(?:anos?|years?|yrs?)"
_EXPERIENCE_WORD = r"(?:experiencia|experience|trayectoria)"
_FILLER = r"(?:profesional|laboral|professional|work|total|comprobable|de|of|over|mas|more|than|con)"

_EXPLICIT_YEARS_RES: tuple[re.Pattern[str], ...] = (
    # "5 años de experiencia", "5+ years of professional experience"
    re.compile(rf"{_NUM}\s*\+?\s*{_YEARS}\s+(?:(?:de|of|en|in)\s+)?(?:\w+\s+){{0,2}}?{_EXPERIENCE_WORD}"),
    # "experiencia de más de 5 años", "experience: 7+ years"
    re.compile(rf"{_EXPERIENCE_WORD}(?:\s*:|\s+{_FILLER})*\s*(?:\+\s*)?{_NUM}\s*\+?\s*{_YEARS}"),
    # "+5 años", "5+ years"
    re.compile(rf"\+\s*{_NUM}\s*{_YEARS}"),
    re.compile(rf"{_NUM}\s*\+\s*{_YEARS}"),
    # "con más de 5 años en...", "with over 5 years in..."
    re.compile(
        rf"\b(?:con|with)\s+(?:(?:mas de|over|more than)\s+|\+\s*)?{_NUM}\s*\+?\s*{_YEARS}"
        r"(?!\s+(?:de edad|old|de vida))"
    ),
)

_MONTHS: dict[int, tuple[str, ...]] = {
    1: ("enero", "ene", "january", "jan"),
    2: ("febrero", "feb", "february"),
    3: ("marzo", "mar", "march"),
    4: ("abril", "abr", "april", "apr"),
    5: ("mayo", "may"),
    6: ("junio", "jun", "june"),
    7: ("julio", "jul", "july"),
    8: ("agosto", "ago", "august", "aug"),
    9: ("septiembre", "setiembre", "september", "sept", "sep", "set"),
    10: ("octubre", "oct", "october"),
    11: ("noviembre", "nov", "november"),
    12: ("diciembre", "dic", "december", "dec"),
}
_MONTH_LOOKUP = {name: number for number, names in _MONTHS.items() for name in names}
_MONTH_NAMES = "|".join(sorted(_MONTH_LOOKUP, key=len, reverse=True))
_PRESENT = (
    r"(?:la actualidad|actualidad|actualmente|actual|presente|present|currently|current|"
    r"today|hoy|now|a la fecha|la fecha)"
)


def _date_pattern(prefix: str) -> str:
    return (
        rf"(?:(?<![a-z])(?P<{prefix}_mname>{_MONTH_NAMES})\.?\s*(?:de\s+|del\s+|,\s*)?"
        rf"|(?<!\d)(?P<{prefix}_mnum>0?[1-9]|1[0-2])\s*[/.-]\s*)?"
        rf"(?P<{prefix}_year>(?:19|20)\d{{2}})"
    )


_DATE_RANGE_RE = re.compile(
    _date_pattern("start")
    + r"\s*(?:-|–|—|a|al|hasta|to|until|till)\s*"
    + rf"(?:{_date_pattern('end')}|(?P<present>{_PRESENT}))(?![\w])"
)

_HEADER_CLEAN_RE = re.compile(r"^[^a-z]+|[^a-z]+$")
_EXPERIENCE_HEADER_RE = re.compile(
    r"(?:experiencia(?: laboral| profesional| relevante)?|historial laboral|"
    r"trayectoria(?: laboral| profesional)?|(?:work|professional|employment|relevant) "
    r"(?:experience|history)|experience|employment)"
)
_EDUCATION_HEADER_RE = re.compile(
    r"(?:educacion|formacion(?: academica)?|estudios|education|academic background|"
    r"certificaciones|certifications|cursos|courses|capacitaciones|training)"
)
_OTHER_HEADER_RE = re.compile(
    r"(?:habilidades|skills|competencias|aptitudes|idiomas|languages|proyectos|projects|"
    r"referencias|references|intereses|interests|voluntariado|volunteer(?:ing)?|premios|awards|"
    r"publicaciones|publications|perfil(?: profesional)?|profile|resumen|summary|sobre mi|"
    r"about me|contacto|contact|datos personales)"
)


def extract_experience_years(text: str, today: date | None = None) -> float | None:
    """Estima los años de experiencia de un CV.

    * Menciones explícitas ("5 años de experiencia", "5+ years", "más de 5 años de
      experiencia", "with 7 years"): se toma la mayor.
    * Rangos de fechas ("Ene 2019 - Actualidad", "2016 - 2020"), fusionando los
      solapados para no contar doble.

    Si el CV tiene un encabezado de experiencia, se suman solo sus rangos y gana el
    mayor de ambos valores ("más de 5 años" es un piso). Sin ese encabezado los rangos
    pueden mezclar estudios, así que solo se usan si no hay mención explícita.

    Args:
        text: Texto del CV.
        today: Fecha de referencia para "Actualidad" (por defecto, hoy).

    Returns:
        Años con un decimal, o ``None`` si no se encontró ninguna señal.
    """
    normalized = normalize_text(text)
    explicit = _explicit_years(normalized)
    lines, has_experience_header = _experience_section_lines(normalized)
    from_ranges = _years_from_date_ranges(lines, today or date.today())
    if explicit is None:
        return from_ranges
    if from_ranges is not None and has_experience_header:
        return max(explicit, from_ranges)
    return explicit


def _explicit_years(normalized: str) -> float | None:
    values: list[float] = []
    for pattern in _EXPLICIT_YEARS_RES:
        for match in pattern.finditer(normalized):
            raw = match.group("n")
            value = float(_NUMBER_WORDS[raw]) if raw in _NUMBER_WORDS else float(raw.replace(",", "."))
            if 0 < value <= MAX_REASONABLE_YEARS:
                values.append(value)
    return max(values) if values else None


def _experience_section_lines(normalized: str) -> tuple[list[str], bool]:
    """Líneas de la sección de experiencia y si se encontró su encabezado.

    Sin encabezado de experiencia, devuelve todo el texto menos la sección de educación.
    """
    section: str | None = None
    found_header = False
    experience: list[str] = []
    not_education: list[str] = []
    for line in normalized.splitlines():
        header = _HEADER_CLEAN_RE.sub("", line).strip(" :")
        if len(header) <= 40:
            if _EXPERIENCE_HEADER_RE.fullmatch(header):
                section = "experience"
                found_header = True
                continue
            if _EDUCATION_HEADER_RE.fullmatch(header):
                section = "education"
                continue
            if _OTHER_HEADER_RE.fullmatch(header):
                section = "other"
                continue
        if section == "experience":
            experience.append(line)
        if section != "education":
            not_education.append(line)
    return (experience, True) if found_header else (not_education, False)


def _years_from_date_ranges(lines: list[str], today: date) -> float | None:
    today_index = today.year * 12 + today.month
    intervals: list[tuple[int, int]] = []
    for match in _DATE_RANGE_RE.finditer("\n".join(lines)):
        start_year = int(match.group("start_year"))
        start_month = _month(match, "start") or 1
        if match.group("present"):
            end_index = today_index
        else:
            end_year = int(match.group("end_year"))
            default_end = 12 if end_year == start_year else 1
            end_index = end_year * 12 + (_month(match, "end") or default_end)
        start_index = start_year * 12 + start_month
        end_index = min(end_index, today_index)
        if start_year < 1960 or start_index > end_index:
            continue
        if (end_index - start_index) / 12 > MAX_REASONABLE_YEARS:
            continue
        intervals.append((start_index, end_index))

    if not intervals:
        return None
    total_months = 0
    current_start, current_end = sorted(intervals)[0]
    for start, end in sorted(intervals)[1:]:
        if start <= current_end + 1:
            current_end = max(current_end, end)
        else:
            total_months += current_end - current_start + 1
            current_start, current_end = start, end
    total_months += current_end - current_start + 1
    return round(total_months / 12, 1)


def _month(match: re.Match[str], prefix: str) -> int | None:
    name = match.group(f"{prefix}_mname")
    if name:
        return _MONTH_LOOKUP[name]
    number = match.group(f"{prefix}_mnum")
    return int(number) if number else None


# ---------------------------------------------------------------------------
# Scoring semántico
# ---------------------------------------------------------------------------


@lru_cache(maxsize=2)
def get_model(model_name: str = DEFAULT_MODEL_NAME) -> SentenceTransformer:
    """Carga el modelo de embeddings una sola vez por proceso.

    Primero intenta desde la caché local, sin tocar la red. Solo si no está lo descarga
    de Hugging Face (~90 MB, una única vez); después todo funciona sin internet.

    Raises:
        RuntimeError: Si falta la librería o el modelo no se puede descargar/cargar.
    """
    # En Windows sin modo desarrollador la caché de Hugging Face no usa symlinks y avisa
    # en cada carga; funciona igual, solo ocupa algo más de disco.
    os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")
    try:
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise RuntimeError("Falta sentence-transformers: ejecutar pip install -r requirements.txt") from exc
    try:
        return SentenceTransformer(model_name, local_files_only=True)
    except Exception:
        logger.info("Modelo '%s' no está en la caché local: descargando...", model_name)
    try:
        return SentenceTransformer(model_name)
    except Exception as exc:
        raise RuntimeError(
            f"No se pudo cargar el modelo '{model_name}'. La primera vez se necesita internet "
            f"para descargarlo de Hugging Face. Detalle: {exc}"
        ) from exc


def _chunk_text(text: str, size: int = CHUNK_WORDS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Parte el texto en ventanas de ``size`` palabras con ``overlap`` de solapamiento."""
    words = text.split()
    if not words:
        return []
    step = max(size - overlap, 1)
    return [" ".join(words[i : i + size]) for i in range(0, max(len(words) - overlap, 1), step)]


def _encode(model: SentenceTransformer, chunks: list[str]) -> np.ndarray:
    return model.encode(
        chunks,
        batch_size=32,
        convert_to_numpy=True,
        normalize_embeddings=True,
        show_progress_bar=False,
    )


@lru_cache(maxsize=8)
def _job_embeddings(model: SentenceTransformer, job_description: str) -> np.ndarray:
    """Embeddings de la JD, cacheados: la misma JD se compara contra todos los CVs."""
    embeddings = _encode(model, _chunk_text(job_description))
    embeddings.setflags(write=False)
    return embeddings


def _calibrate(similarity: float) -> float:
    span = SIMILARITY_CEILING - SIMILARITY_FLOOR
    return float(np.clip((similarity - SIMILARITY_FLOOR) / span, 0.0, 1.0) * 100.0)


def score_cv(cv_text: str, job_description: str, model: SentenceTransformer | None = None) -> float:
    """Score semántico 0-100 entre un CV y una Job Description.

    Ambos textos se parten en ventanas (el modelo trunca a 256 tokens). Para cada
    fragmento de la JD se toma el fragmento del CV más parecido (cosine similarity)
    y se promedian esos máximos: mide cuánto de la JD está cubierto por el CV.
    El coseno resultante se escala de [SIMILARITY_FLOOR, SIMILARITY_CEILING] a [0, 100].

    Args:
        cv_text: Texto del CV.
        job_description: Texto de la búsqueda.
        model: Modelo ya cargado; si es ``None`` se usa :func:`get_model`.

    Raises:
        ValueError: Si la Job Description está vacía.
    """
    job_description = " ".join(job_description.split())
    if not job_description:
        raise ValueError("La Job Description está vacía.")
    cv_chunks = _chunk_text(cv_text)
    if not cv_chunks:
        return 0.0
    if model is None:
        model = get_model()

    similarity = cosine_similarity(_job_embeddings(model, job_description), _encode(model, cv_chunks))
    return round(_calibrate(float(similarity.max(axis=1).mean())), 1)


def combine_scores(semantic_score: float, coverage: float | None, semantic_weight: float) -> float:
    """Score final = peso * semántico + (1 - peso) * cobertura de skills.

    Sin skills requeridas (``coverage is None``) el score final es el semántico.
    """
    if coverage is None:
        return round(semantic_score, 1)
    weight = min(max(semantic_weight, 0.0), 1.0)
    return round(weight * semantic_score + (1.0 - weight) * coverage, 1)


def get_verdict(score: float) -> str:
    """PASA (> 65), REVISAR (45-65) o NO PASA (< 45)."""
    if score > PASS_THRESHOLD:
        return VERDICT_PASS
    if score >= REVIEW_THRESHOLD:
        return VERDICT_REVIEW
    return VERDICT_FAIL


# ---------------------------------------------------------------------------
# Orquestación
# ---------------------------------------------------------------------------


def analyze_cv(
    file_name: str,
    cv_text: str,
    job_description: str,
    required_skills: Sequence[str],
    *,
    model: SentenceTransformer | None = None,
    semantic_weight: float = DEFAULT_SEMANTIC_WEIGHT,
    warnings: Iterable[str] = (),
) -> CandidateResult:
    """Evalúa un CV ya extraído y devuelve todas las métricas del ranking."""
    required = [canonicalize_skill(skill) for skill in required_skills if skill.strip()]
    parsed = parse_cv(cv_text, extra_skills=required)
    semantic = score_cv(cv_text, job_description, model=model)
    coverage = skill_coverage(parsed.skills, required)
    score = combine_scores(semantic, coverage, semantic_weight)

    required_found = [skill for skill in dict.fromkeys(required) if skill in parsed.skills]
    other_found = [skill for skill in parsed.skills if skill not in required_found]
    return CandidateResult(
        file_name=file_name,
        email=parsed.email,
        phone=parsed.phone,
        linkedin=parsed.linkedin,
        score=score,
        semantic_score=semantic,
        skill_coverage=coverage,
        found_skills=required_found + other_found,
        missing_skills=get_missing_skills(parsed.skills, required),
        experience_years=extract_experience_years(cv_text),
        verdict=get_verdict(score),
        warnings=list(warnings),
    )


def analyze_files(
    files: Sequence[tuple[str, bytes]],
    job_description: str,
    required_skills: Sequence[str],
    *,
    model: SentenceTransformer | None = None,
    semantic_weight: float = DEFAULT_SEMANTIC_WEIGHT,
    extract: Callable[[bytes, str], Extraction] = extract_from_bytes,
    on_progress: Callable[[int, int, str], None] | None = None,
) -> tuple[list[CandidateResult], list[FileFailure]]:
    """Evalúa un lote de CVs. Un archivo fallido se registra y no frena el resto.

    Args:
        files: Pares ``(nombre, bytes)``.
        job_description: Texto de la búsqueda.
        required_skills: Skills requeridas (se canonizan).
        model: Modelo de embeddings ya cargado (opcional).
        semantic_weight: Peso del score semántico en el score final (0-1).
        extract: Función de extracción; permite inyectar una versión cacheada.
        on_progress: Callback ``(procesados, total, nombre_actual)``.

    Returns:
        Resultados de los CVs evaluados y lista de archivos que fallaron.
    """
    if model is None:
        model = get_model()
    results: list[CandidateResult] = []
    failures: list[FileFailure] = []
    total = len(files)
    for index, (name, data) in enumerate(files, start=1):
        if on_progress:
            on_progress(index - 1, total, name)
        try:
            extraction = extract(data, name)
            results.append(
                analyze_cv(
                    name,
                    extraction.text,
                    job_description,
                    required_skills,
                    model=model,
                    semantic_weight=semantic_weight,
                    warnings=extraction.warnings,
                )
            )
        except ScannedPDFError as exc:
            failures.append(FileFailure(name, "PDF escaneado", str(exc)))
        except UnsupportedFileError as exc:
            failures.append(FileFailure(name, "Formato no soportado", str(exc)))
        except ExtractionError as exc:
            failures.append(FileFailure(name, "Error de lectura", str(exc)))
        except Exception as exc:  # un archivo problemático no debe tumbar el lote
            logger.exception("Error inesperado procesando %s", name)
            failures.append(FileFailure(name, "Error inesperado", f"{type(exc).__name__}: {exc}"))
    if on_progress:
        on_progress(total, total, "")
    return results, failures


# ---------------------------------------------------------------------------
# Tablas y exportación
# ---------------------------------------------------------------------------

DISPLAY_COLUMNS: list[str] = [
    "Archivo", "Email", "Score %", "Skills Encontradas", "Skills Faltantes", "Años Exp", "Veredicto",
]
EXPORT_COLUMNS: list[str] = [
    "Archivo", "Email", "Teléfono", "LinkedIn", "Score %", "Score Semántico %", "Cobertura Skills %",
    "Skills Encontradas", "Skills Faltantes", "Años Exp", "Veredicto", "Advertencias",
]
FAILURE_COLUMNS: list[str] = ["Archivo", "Tipo", "Detalle"]

_VERDICT_FILLS = {
    VERDICT_PASS: ("C6EFCE", "006100"),
    VERDICT_REVIEW: ("FFEB9C", "9C5700"),
    VERDICT_FAIL: ("FFC7CE", "9C0006"),
}


def results_to_dataframe(results: Iterable[CandidateResult]) -> pd.DataFrame:
    """Tabla de ranking ordenada por score descendente (columnas de :data:`EXPORT_COLUMNS`)."""
    rows = [
        {
            "Archivo": r.file_name,
            "Email": r.email,
            "Teléfono": r.phone,
            "LinkedIn": r.linkedin,
            "Score %": r.score,
            "Score Semántico %": r.semantic_score,
            "Cobertura Skills %": r.skill_coverage,
            "Skills Encontradas": ", ".join(r.found_skills),
            "Skills Faltantes": ", ".join(r.missing_skills),
            "Años Exp": r.experience_years,
            "Veredicto": r.verdict,
            "Advertencias": " | ".join(r.warnings),
        }
        for r in results
    ]
    frame = pd.DataFrame(rows, columns=EXPORT_COLUMNS)
    frame = frame.sort_values(["Score %", "Score Semántico %"], ascending=False, kind="stable")
    return frame.reset_index(drop=True)


def failures_to_dataframe(failures: Iterable[FileFailure]) -> pd.DataFrame:
    """Tabla de archivos no procesados."""
    rows = [{"Archivo": f.file_name, "Tipo": f.kind, "Detalle": f.detail} for f in failures]
    return pd.DataFrame(rows, columns=FAILURE_COLUMNS)


def to_csv_bytes(ranking: pd.DataFrame) -> bytes:
    """CSV en UTF-8 con BOM para que Excel respete tildes y eñes."""
    return ranking.to_csv(index=False).encode("utf-8-sig")


def to_excel_bytes(
    ranking: pd.DataFrame,
    failures: pd.DataFrame | None = None,
    parameters: dict[str, Any] | None = None,
) -> bytes:
    """Excel con hojas Ranking (con formato y colores por veredicto), Errores y Parámetros."""
    from openpyxl.styles import Alignment, Font, PatternFill

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        sheets: dict[str, pd.DataFrame] = {"Ranking": ranking}
        if failures is not None and not failures.empty:
            sheets["Errores"] = failures
        if parameters:
            sheets["Parámetros"] = pd.DataFrame(
                {"Parámetro": list(parameters), "Valor": [str(v) for v in parameters.values()]}
            )

        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            sheet = writer.sheets[sheet_name]
            sheet.freeze_panes = "A2"
            if sheet.max_row > 1:
                sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for index, column in enumerate(frame.columns, start=1):
                longest = max([len(str(column)), *(len(str(v)) for v in frame[column].tolist())])
                sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(longest + 2, 10), 70)

        ranking_sheet = writer.sheets["Ranking"]
        columns = list(ranking.columns)
        for name in ("Score %", "Score Semántico %", "Cobertura Skills %", "Años Exp"):
            if name in columns:
                for row in range(2, ranking_sheet.max_row + 1):
                    ranking_sheet.cell(row, columns.index(name) + 1).number_format = "0.0"
        if "Veredicto" in columns:
            verdict_col = columns.index("Veredicto") + 1
            for row in range(2, ranking_sheet.max_row + 1):
                cell = ranking_sheet.cell(row, verdict_col)
                fill, font = _VERDICT_FILLS.get(str(cell.value), (None, None))
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                    cell.font = Font(bold=True, color=font)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def _read_job_description(value: str) -> str:
    path = Path(value)
    if path.suffix.lower() in {".txt", ".md"} and path.is_file():
        return path.read_text(encoding="utf-8")
    return value


def main(argv: Sequence[str] | None = None) -> int:
    """Procesa una carpeta de CVs desde la terminal y guarda el ranking en CSV o Excel."""
    parser = argparse.ArgumentParser(description="Filtro ATS local: ranking de CVs contra una Job Description.")
    parser.add_argument("folder", help="Carpeta con CVs en PDF/DOCX")
    parser.add_argument("--jd", required=True, help="Job Description: texto o ruta a un .txt/.md")
    parser.add_argument("--skills", default="", help='Skills requeridas separadas por coma, ej. "python, sql"')
    parser.add_argument("--weight", type=float, default=DEFAULT_SEMANTIC_WEIGHT, help="Peso semántico 0-1 (default 0.7)")
    parser.add_argument("--out", default="ranking_ats.xlsx", help="Archivo de salida .xlsx o .csv")
    parser.add_argument("--recursive", action="store_true", help="Incluir subcarpetas")
    args = parser.parse_args(argv)

    logging.basicConfig(level=logging.WARNING, format="%(levelname)s %(message)s")
    logger.setLevel(logging.INFO)
    job_description = _read_job_description(args.jd)
    if not job_description.strip():
        parser.error("La Job Description está vacía.")
    try:
        paths = iter_cv_files(args.folder, recursive=args.recursive)
    except NotADirectoryError as exc:
        parser.error(str(exc))
    if not paths:
        parser.error(f"No hay archivos PDF/DOCX en {args.folder}")

    root = Path(args.folder)
    files = [(str(path.relative_to(root)), path.read_bytes()) for path in paths]
    required = parse_skill_list(args.skills)
    logger.info("Analizando %d CVs con el modelo %s...", len(files), DEFAULT_MODEL_NAME)
    results, failures = analyze_files(files, job_description, required, semantic_weight=args.weight)

    ranking = results_to_dataframe(results)
    failures_df = failures_to_dataframe(failures)
    out = Path(args.out)
    if out.suffix.lower() == ".csv":
        out.write_bytes(to_csv_bytes(ranking))
    else:
        parameters = {"Skills requeridas": ", ".join(required), "Peso semántico": args.weight, "Job Description": job_description}
        out.write_bytes(to_excel_bytes(ranking, failures_df, parameters))

    with pd.option_context("display.max_columns", None, "display.width", 200, "display.max_colwidth", 40):
        print(ranking[DISPLAY_COLUMNS].to_string(index=False))
    for failure in failures:
        logger.warning("%s -> %s: %s", failure.file_name, failure.kind, failure.detail)
    logger.info("Ranking guardado en %s", out.resolve())
    return 0


if __name__ == "__main__":
    sys.exit(main())
